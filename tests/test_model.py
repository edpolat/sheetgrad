import pytest
from openpyxl import Workbook

from sheetgrad.model import CircularReferenceError, FormulaError, Model


def excel_yap(tmp_path, hucreler):
    wb = Workbook()
    ws = wb.active
    for ref, deger in hucreler.items():
        ws[ref] = deger
    yol = tmp_path / "model.xlsx"
    wb.save(yol)
    return Model.from_excel(yol)


@pytest.fixture
def kar_modeli(tmp_path):
    return excel_yap(tmp_path, {
        "A2": "Fiyat", "B2": 100,
        "A3": "Birim maliyet", "B3": 60,
        "A4": "Reklam bütçesi", "B4": 20000,
        "A5": "Kira", "B5": 15000,
        "A6": "Maaşlar", "B6": 40000,
        "A9": "Satış adedi", "B9": "=5000-40*B2+0.05*B4",
        "A10": "Ciro", "B10": "=B2*B9",
        "A11": "Değişken maliyet", "B11": "=B3*B9",
        "A12": "Toplam gider", "B12": "=B11+B4+B5+B6",
        "A13": "Kâr", "B13": "=B10-B12",
    })


def test_hesaplanan_degerler(kar_modeli):
    assert kar_modeli.cell("B9").data == 2000
    assert kar_modeli.cell("B13").data == 5000


def test_duyarliliklar(kar_modeli):
    kar, satirlar = kar_modeli.sensitivities("B13")
    assert kar == 5000
    turevler = {s["label"]: s["gradient"] for s in satirlar}
    assert turevler == pytest.approx({
        "Fiyat": 400,              # adet + (fiyat - maliyet) × (-40)
        "Birim maliyet": -2000,    # -adet
        "Reklam bütçesi": 1,       # (fiyat - maliyet) × 0.05 - 1
        "Kira": -1,
        "Maaşlar": -1,
    })


def test_yuzde_bir_etkisi(kar_modeli):
    _, satirlar = kar_modeli.sensitivities("B13")
    etkiler = {s["label"]: s["impact_1pct"] for s in satirlar}
    assert etkiler == pytest.approx({
        "Birim maliyet": -1200,    # -2000 × 60 × %1
        "Fiyat": 400,              # 400 × 100 × %1
        "Maaşlar": -400,
        "Reklam bütçesi": 200,     # 1 × 20000 × %1
        "Kira": -150,
    })


def test_yuzde_etkiye_gore_sirali(kar_modeli):
    _, satirlar = kar_modeli.sensitivities("B13")
    mutlak = [abs(s["impact_1pct"]) for s in satirlar]
    assert mutlak == sorted(mutlak, reverse=True)
    assert satirlar[0]["label"] == "Birim maliyet"


def test_iki_kez_cagrilinca_turevler_birikmez(kar_modeli):
    _, ilk = kar_modeli.sensitivities("B13")
    _, ikinci = kar_modeli.sensitivities("B13")
    assert ilk == ikinci


def test_sadece_kullanilan_girdiler(tmp_path):
    model = excel_yap(tmp_path, {"B1": 2, "B2": 3, "B3": 99, "B4": "=B1*B2"})
    _, satirlar = model.sensitivities("B4")
    assert {s["cell"] for s in satirlar} == {"B1", "B2"}


def test_bos_hucre_sifir_sayilir(tmp_path):
    model = excel_yap(tmp_path, {"B1": 5, "B2": "=B1+B7"})
    assert model.cell("B2").data == 5


def test_dongusel_basvuru(tmp_path):
    model = excel_yap(tmp_path, {"A1": "=B1+1", "B1": "=C1*2", "C1": "=A1"})
    with pytest.raises(CircularReferenceError, match="A1 → B1 → C1 → A1"):
        model.cell("A1")


def test_metin_hucresi_hata_verir(tmp_path):
    model = excel_yap(tmp_path, {"A1": "merhaba", "B1": "=A1*2"})
    with pytest.raises(FormulaError, match="A1"):
        model.cell("B1")


def test_hata_mesaji_hucreyi_soyler(tmp_path):
    model = excel_yap(tmp_path, {"A1": 1, "B1": "=VLOOKUP(A1, A1:A1, 1)"})
    with pytest.raises(FormulaError, match="B1.*VLOOKUP"):
        model.cell("B1")


def test_baska_sayfa_desteklenmiyor(tmp_path):
    model = excel_yap(tmp_path, {"B1": "=Sayfa2!A1*2"})
    with pytest.raises(FormulaError, match="B1"):
        model.cell("B1")


def test_etiket_yoksa_hucre_adi(tmp_path):
    model = excel_yap(tmp_path, {"B1": 2, "B2": "=B1*3"})
    _, satirlar = model.sensitivities("B2")
    assert satirlar[0]["label"] == "B1"
