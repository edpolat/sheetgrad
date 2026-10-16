import pytest
from openpyxl import Workbook

from sheetgrad.engine import Value
from sheetgrad.formula import evaluate, excel_round, tokenize
from sheetgrad.model import Model


def hucreler(**degerler):
    tablo = {ad: Value(float(v)) for ad, v in degerler.items()}
    return tablo, lambda ref: tablo[ref]


def hesapla(formul, **degerler):
    tablo, get = hucreler(**degerler)
    notlar = []
    sonuc = evaluate(formul, get, notlar)
    sonuc.backward()
    return sonuc, tablo, notlar


# --- karşılaştırma ---

def test_karsilastirma_tokenleri():
    assert tokenize("=A1<=B1") == [("CELL", "A1"), ("OP", "<="), ("CELL", "B1")]
    assert tokenize("=A1<>B1")[1] == ("OP", "<>")


@pytest.mark.parametrize("formul, beklenen", [
    ("=3>2", 1), ("=3<2", 0), ("=2=2", 1), ("=2<>2", 0),
    ("=2>=2", 1), ("=1+1=2", 1),   # karşılaştırma en düşük öncelikte
])
def test_karsilastirma(formul, beklenen):
    sonuc, _, _ = hesapla(formul)
    assert sonuc.data == beklenen


# --- IF ---

def test_if_dogru_dal_ve_turevi():
    sonuc, t, _ = hesapla("=IF(A1>100, A1*2, A1*3)", A1=150)
    assert sonuc.data == 300
    assert t["A1"].grad == 2


def test_if_yanlis_dal_ve_turevi():
    sonuc, t, _ = hesapla("=IF(A1>100, A1*2, A1*3)", A1=50)
    assert sonuc.data == 150
    assert t["A1"].grad == 3


def test_if_secilmeyen_dal_hesaplanmaz():
    # Excel'deki yaygın kalıp: sıfıra bölmeyi IF ile önlemek
    sonuc, _, _ = hesapla("=IF(A1=0, 0, B1/A1)", A1=0, B1=10)
    assert sonuc.data == 0
    sonuc, _, _ = hesapla("=IF(A1>0, LN(A1), 0)", A1=-5)
    assert sonuc.data == 0


def test_if_secilmeyen_daldaki_hucre_okunmaz():
    okunan = []
    tablo = {"A1": Value(1.0), "B1": Value(5.0), "C1": Value(7.0)}

    def get(ref):
        okunan.append(ref)
        return tablo[ref]

    evaluate("=IF(A1>0, B1, C1)", get)
    assert okunan == ["A1", "B1"]


def test_if_ic_ice_parantezli_dal_atlanir():
    sonuc, _, _ = hesapla("=IF(A1>0, 1, SUM(A1, (A1+1)*2))", A1=1)
    assert sonuc.data == 1


def test_if_iki_argumanli():
    sonuc, _, _ = hesapla("=IF(A1>0, 5)", A1=-1)
    assert sonuc.data == 0


def test_if_sinira_yakin_uyari():
    _, _, notlar = hesapla("=IF(A1>100, 1, 2)", A1=100.5)
    assert any("sınıra çok yakın" in n for n in notlar)
    _, _, notlar = hesapla("=IF(A1>100, 1, 2)", A1=150)
    assert notlar == []


# --- MIN / MAX / ABS ---

def test_max_turevi_secilene_gider():
    sonuc, t, notlar = hesapla("=MAX(A1, A2, A3)", A1=3, A2=7, A3=5)
    assert sonuc.data == 7
    assert (t["A1"].grad, t["A2"].grad, t["A3"].grad) == (0, 1, 0)
    assert notlar == []


def test_min_aralik_ile():
    sonuc, t, _ = hesapla("=MIN(A1:A3)*10", A1=3, A2=7, A3=5)
    assert sonuc.data == 30
    assert t["A1"].grad == 10


def test_min_esitlik_uyarisi():
    _, _, notlar = hesapla("=MIN(A1, A2)", A1=4, A2=4)
    assert any("eşit değerler" in n for n in notlar)


def test_vergi_formulu():
    # Kâr negatifse vergi 0, pozitifse %20
    sonuc, t, _ = hesapla("=MAX(0, A1)*0.2", A1=1000)
    assert sonuc.data == 200
    assert t["A1"].grad == pytest.approx(0.2)
    sonuc, t, _ = hesapla("=MAX(0, A1)*0.2", A1=-1000)
    assert sonuc.data == 0
    assert t["A1"].grad == 0


@pytest.mark.parametrize("x, deger, turev", [(-3, 3, -1), (3, 3, 1)])
def test_abs(x, deger, turev):
    sonuc, t, _ = hesapla("=ABS(A1)", A1=x)
    assert sonuc.data == deger
    assert t["A1"].grad == turev


def test_abs_sifirda_uyari():
    _, t, notlar = hesapla("=ABS(A1)", A1=0)
    assert t["A1"].grad == 0
    assert notlar


# --- ROUND ---

@pytest.mark.parametrize("sayi, basamak, beklenen", [
    (2.5, 0, 3), (-2.5, 0, -3), (3.14159, 2, 3.14), (1234, -2, 1200), (0.5, 0, 1),
])
def test_excel_round(sayi, basamak, beklenen):
    # Python'un round(2.5) = 2 sonucundan farklı olarak Excel 3 verir
    assert excel_round(sayi, basamak) == pytest.approx(beklenen)


def test_round_turevi_gecirir():
    sonuc, t, notlar = hesapla("=ROUND(A1*A2, 0)", A1=3.3, A2=10)
    assert sonuc.data == 33
    assert t["A1"].grad == pytest.approx(10)
    assert any("ROUND" in n for n in notlar)


# --- modele entegrasyon ---

def test_model_uyarilari_hucre_adiyla(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["B1"] = 100.5
    ws["B2"] = "=IF(B1>100, B1*2, 0)"
    yol = tmp_path / "m.xlsx"
    wb.save(yol)
    model = Model.from_excel(yol)
    model.sensitivities("B2")
    assert any(n.startswith("B2:") for n in model.notes)
