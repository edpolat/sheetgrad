from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "Kar Modeli"

satirlar = [
    ("GİRDİLER", None),
    ("Fiyat", 100),
    ("Birim maliyet", 60),
    ("Reklam bütçesi", 20000),
    ("Kira", 15000),
    ("Maaşlar", 40000),
    (None, None),
    ("HESAPLAR", None),
    ("Satış adedi", "=5000-40*B2+0.05*B4"),
    ("Ciro", "=B2*B9"),
    ("Değişken maliyet", "=B3*B9"),
    ("Toplam gider", "=B11+B4+B5+B6"),
    ("Kâr", "=B10-B12"),
]

for satir, (etiket, deger) in enumerate(satirlar, start=1):
    ws.cell(row=satir, column=1, value=etiket)
    ws.cell(row=satir, column=2, value=deger)

ws.column_dimensions["A"].width = 20
wb.save("examples/kar_modeli.xlsx")
print("examples/kar_modeli.xlsx oluşturuldu")
