import openpyxl

from sheetgrad.engine import Value
from sheetgrad.formula import evaluate


class CircularReferenceError(Exception):
    pass


class FormulaError(Exception):
    pass


class Model:
    def __init__(self, sheet):
        self.sheet = sheet
        self._cache = {}
        self._in_progress = []
        self.inputs = {}

    @classmethod
    def from_excel(cls, path, sheet_name=None):
        workbook = openpyxl.load_workbook(path, data_only=False)
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        return cls(sheet)

    def reset(self):
        self._cache = {}
        self._in_progress = []
        self.inputs = {}

    def cell(self, ref):
        ref = ref.upper()
        if ref in self._cache:
            return self._cache[ref]

        if ref in self._in_progress:
            chain = " → ".join(self._in_progress[self._in_progress.index(ref):] + [ref])
            raise CircularReferenceError(f"Döngüsel başvuru: {chain}")

        raw = self.sheet[ref].value
        self._in_progress.append(ref)
        try:
            value = self._to_value(ref, raw)
        finally:
            self._in_progress.pop()

        self._cache[ref] = value
        return value

    def _to_value(self, ref, raw):
        if raw is None:
            return Value(0.0)

        if isinstance(raw, bool):
            raise FormulaError(f"{ref} hücresinde mantıksal değer var: {raw}")

        if isinstance(raw, (int, float)):
            value = Value(float(raw))
            self.inputs[ref] = value
            return value

        if isinstance(raw, str) and raw.startswith("="):
            try:
                return evaluate(raw, self.cell)
            except (CircularReferenceError, FormulaError):
                raise
            except Exception as error:
                raise FormulaError(f"{ref} hücresinde ({raw}): {error}") from error

        raise FormulaError(f"{ref} hücresinde sayı yerine metin var: {raw!r}")

    def label(self, ref):
        column = ref.rstrip("0123456789")
        row = ref[len(column):]
        if column == "A":
            return ref
        left = self.sheet.cell(row=int(row), column=self.sheet[ref].column - 1).value
        return left if isinstance(left, str) and left.strip() else ref

    def sensitivities(self, output_ref):
        self.reset()
        output = self.cell(output_ref)
        output.backward()
        rows = [
            {
                "cell": ref,
                "label": self.label(ref),
                "value": value.data,
                "gradient": value.grad,
            }
            for ref, value in self.inputs.items()
        ]
        rows.sort(key=lambda row: abs(row["gradient"]), reverse=True)
        return output.data, rows
